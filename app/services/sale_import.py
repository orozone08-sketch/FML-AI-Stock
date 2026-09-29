from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import itertools
import re

from openpyxl import load_workbook

from app.extensions import db
from app.models import Company, Customer, Item, Sale, StockBook, normalize_customer_name
from app.services.transactions import create_sale


HEADER_NAMES = {
    "date": "Date",
    "particulars": "Particulars",
    "voucher_type": "Voucher Type",
    "voucher_number": "Voucher No.",
    "payment_terms": "Terms of Payment",
    "delivery_terms": "Terms of Delivery",
    "quantity": "Quantity",
    "rate": "Rate",
    "value": "Value",
    "gross_total": "Gross Total",
    "igst": "Output IGST @ 18%",
    "cgst": "Output CGST @ 9%",
    "sgst": "Output SGST @ 9%",
    "round_off": "Round Off",
}


def _text(value):
    return " ".join(str(value or "").strip().split())


def _decimal(value, default=Decimal("0.00")):
    if value in (None, ""):
        return default
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return default


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _header_map(rows):
    for index, row in enumerate(rows[:30]):
        values = {_text(value).casefold(): column for column, value in enumerate(row)}
        if {"date", "particulars", "voucher no."}.issubset(values):
            return index, {
                key: values.get(label.casefold())
                for key, label in HEADER_NAMES.items()
            }
    raise ValueError("Could not find a Sales Register header row.")


def _value(row, columns, key):
    column = columns.get(key)
    return row[column] if column is not None and column < len(row) else None


def parse_sales_workbook(file_stream, filename="workbook"):
    try:
        workbook = load_workbook(file_stream, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("This file is not a supported Excel workbook.") from exc

    if not workbook.sheetnames:
        raise ValueError("The workbook has no worksheets.")
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    header_index, columns = _header_map(rows)
    invoices = []
    current = None

    for row in rows[header_index + 1 :]:
        particulars = _text(_value(row, columns, "particulars"))
        if not particulars:
            continue
        if particulars.casefold() in {"grand total", "total"}:
            continue

        voucher_number = _text(_value(row, columns, "voucher_number"))
        voucher_type = _text(_value(row, columns, "voucher_type"))
        invoice_date = _date(_value(row, columns, "date"))
        is_invoice = bool(voucher_number and (invoice_date or voucher_type))

        if is_invoice:
            gst = sum(
                (_decimal(_value(row, columns, key)) for key in ("igst", "cgst", "sgst")),
                Decimal("0.00"),
            )
            subtotal = _decimal(_value(row, columns, "value"))
            grand_total = _decimal(_value(row, columns, "gross_total"), subtotal + gst)
            current = {
                "date": invoice_date,
                "customer": particulars,
                "invoice_number": voucher_number,
                "payment_terms": _text(_value(row, columns, "payment_terms")),
                "delivery_terms": _text(_value(row, columns, "delivery_terms")),
                "subtotal": subtotal,
                "gst_total": gst,
                "grand_total": grand_total,
                "round_off": _decimal(_value(row, columns, "round_off")),
                "lines": [],
            }
            invoices.append(current)
            continue

        if current is None:
            continue
        quantity = _decimal(_value(row, columns, "quantity"), None)
        rate = _decimal(_value(row, columns, "rate"), None)
        value = _decimal(_value(row, columns, "value"), None)
        if quantity is None and rate is None and value is None:
            continue
        current["lines"].append(
            {
                "item": particulars,
                "quantity": quantity or Decimal("0.000"),
                "rate": rate or Decimal("0.00"),
                "value": value or Decimal("0.00"),
            }
        )

    if not invoices:
        raise ValueError("No sales invoices were found in the workbook.")
    return {
        "filename": filename,
        "sheet": sheet.title,
        "invoices": invoices,
        "invoice_count": len(invoices),
        "line_count": sum(len(invoice["lines"]) for invoice in invoices),
    }


def _unique_code(model, prefix, name):
    slug = re.sub(r"[^A-Z0-9]+", "-", _text(name).upper()).strip("-") or "RECORD"
    base = f"{prefix}-{slug}"[:45]
    for number in itertools.count(1):
        suffix = "" if number == 1 else f"-{number}"
        candidate = f"{base[:50 - len(suffix)]}{suffix}"
        if not db.session.query(model.id).filter(model.code == candidate).first():
            return candidate


def _customer_for_import(name, user):
    name = _text(name)
    key = normalize_customer_name(name)
    customer = Customer.query.filter_by(name_key=key).first()
    if customer:
        return customer, False
    customer = Customer(
        code=_unique_code(Customer, "IMP-CUST", name),
        name=name,
        name_key=key,
        customer_type="CASH_AND_BILL",
        default_credit_days=30,
        active=True,
        created_by_id=getattr(user, "id", None),
    )
    db.session.add(customer)
    db.session.flush()
    return customer, True


def _item_for_import(name, gst_percent, user):
    name = _text(name)
    item = Item.query.filter(db.func.lower(db.func.trim(Item.name)) == name.casefold()).first()
    if item:
        return item, False
    item = Item(
        code=_unique_code(Item, "IMP-ITEM", name),
        name=name,
        unit="pcs",
        gst_percent=gst_percent,
        minimum_stock=Decimal("0.000"),
        active=True,
        created_by_id=getattr(user, "id", None),
    )
    db.session.add(item)
    db.session.flush()
    return item, True


def import_sales_workbook(preview, company_id, stock_book_id, user):
    company = db.session.get(Company, int(company_id or 0))
    stock_book = db.session.get(StockBook, int(stock_book_id or 0))
    if not company or not company.active:
        raise ValueError("Select an active company.")
    if not stock_book or not stock_book.active or stock_book.company_id != company.id:
        raise ValueError("Select an active stock book belonging to the selected company.")
    if stock_book.book_type != "GST":
        raise ValueError("This sales register contains GST columns, so select a GST stock book.")

    result = {
        "created": [],
        "skipped": [],
        "created_customers": 0,
        "created_items": 0,
    }
    for invoice in preview["invoices"]:
        existing = Sale.query.filter_by(
            company_id=company.id,
            invoice_number=invoice["invoice_number"],
        ).first()
        if existing:
            result["skipped"].append(
                {"invoice_number": invoice["invoice_number"], "reason": "Invoice already exists."}
            )
            continue

        invalid_lines = [
            line["item"]
            for line in invoice["lines"]
            if line["quantity"] <= 0 or line["rate"] <= 0
        ]
        if invalid_lines:
            result["skipped"].append(
                {
                    "invoice_number": invoice["invoice_number"],
                    "reason": "Missing quantity or rate for: " + ", ".join(invalid_lines),
                }
            )
            continue
        if not invoice["date"] or not invoice["lines"]:
            result["skipped"].append(
                {"invoice_number": invoice["invoice_number"], "reason": "Missing date or product lines."}
            )
            continue

        try:
            with db.session.begin_nested():
                customer, customer_created = _customer_for_import(invoice["customer"], user)
                line_rows = []
                gst_percent = Decimal("18.00") if invoice["gst_total"] > 0 else Decimal("0.00")
                item_created_count = 0
                for line in invoice["lines"]:
                    item, item_created = _item_for_import(line["item"], gst_percent, user)
                    item_created_count += int(item_created)
                    line_rows.append(
                        {
                            "item_id": item.id,
                            "quantity": line["quantity"],
                            "rate": line["rate"],
                            "gst_percent": gst_percent,
                        }
                    )
                remarks = f"Imported from {preview['filename']}"
                if invoice["payment_terms"]:
                    remarks += f". Payment terms: {invoice['payment_terms']}"
                sale = create_sale(
                    {
                        "company_id": company.id,
                        "stock_book_id": stock_book.id,
                        "customer_id": customer.id,
                        "sale_type": "GST",
                        "invoice_number": invoice["invoice_number"],
                        "invoice_date": invoice["date"].isoformat(),
                        "remarks": remarks,
                    },
                    line_rows,
                    user,
                )
            result["created"].append(sale.invoice_number)
            result["created_customers"] += int(customer_created)
            result["created_items"] += item_created_count
        except Exception as exc:
            result["skipped"].append(
                {"invoice_number": invoice["invoice_number"], "reason": str(exc)}
            )
    return result