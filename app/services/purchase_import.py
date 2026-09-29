from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
import itertools
import re

from openpyxl import load_workbook

from app.extensions import db
from app.models import Company, Item, Purchase, StockBook, Supplier
from app.services.transactions import create_purchase


HEADERS = {
    "date": "Date", "particulars": "Particulars", "voucher_number": "Voucher No.",
    "voucher_type": "Voucher Type", "payment_terms": "Terms of Payment",
    "quantity": "Quantity", "rate": "Rate", "value": "Value",
    "gross_total": "Gross Total", "gst_net": "Purchase GST Net @ 18%",
    "cgst": "Input CGST@ 9%", "sgst": "Input SGST @ 9%", "round_off": "Round Off",
}


def _text(value):
    return " ".join(str(value or "").strip().split())


def _decimal(value, default=Decimal("0.00")):
    if value in (None, ""):
        return default
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return default


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    return value if isinstance(value, date) else None


def _columns(rows):
    for index, row in enumerate(rows[:30]):
        found = {_text(value).casefold(): i for i, value in enumerate(row)}
        if {"date", "particulars", "voucher no."}.issubset(found):
            return index, {key: found.get(label.casefold()) for key, label in HEADERS.items()}
    raise ValueError("Could not find a Purchase Register header row.")


def _cell(row, columns, key):
    index = columns.get(key)
    return row[index] if index is not None and index < len(row) else None


def parse_purchase_workbook(file_stream, filename="workbook"):
    try:
        workbook = load_workbook(file_stream, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("This file is not a supported Excel workbook.") from exc
    if not workbook.sheetnames:
        raise ValueError("The workbook has no worksheets.")
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    header, columns = _columns(rows)
    purchases, current = [], None
    for row in rows[header + 1:]:
        name = _text(_cell(row, columns, "particulars"))
        if not name or name.casefold() in {"grand total", "total"}:
            continue
        number = _text(_cell(row, columns, "voucher_number"))
        bill_date = _date(_cell(row, columns, "date"))
        if number and (bill_date or _text(_cell(row, columns, "voucher_type"))):
            gst_net = _decimal(_cell(row, columns, "gst_net"))
            cgst = _decimal(_cell(row, columns, "cgst"))
            sgst = _decimal(_cell(row, columns, "sgst"))
            current = {
                "date": bill_date, "supplier": name, "bill_number": number,
                "payment_terms": _text(_cell(row, columns, "payment_terms")),
                "subtotal": gst_net, "gst_total": cgst + sgst,
                "grand_total": _decimal(_cell(row, columns, "gross_total"), gst_net + cgst + sgst),
                "round_off": _decimal(_cell(row, columns, "round_off")), "lines": [],
            }
            purchases.append(current)
            continue
        if current is None:
            continue
        quantity = _decimal(_cell(row, columns, "quantity"), None)
        rate = _decimal(_cell(row, columns, "rate"), None)
        value = _decimal(_cell(row, columns, "value"), None)
        if quantity is not None or rate is not None or value is not None:
            current["lines"].append({
                "item": name, "quantity": quantity or Decimal("0.000"),
                "rate": rate or Decimal("0.00"), "value": value or Decimal("0.00"),
            })
    if not purchases:
        raise ValueError("No purchase bills were found in the workbook.")
    return {
        "filename": filename, "sheet": sheet.title, "purchases": purchases,
        "purchase_count": len(purchases),
        "line_count": sum(len(p["lines"]) for p in purchases),
    }


def _unique_code(model, prefix, name):
    slug = re.sub(r"[^A-Z0-9]+", "-", _text(name).upper()).strip("-") or "RECORD"
    base = f"{prefix}-{slug}"[:45]
    for number in itertools.count(1):
        suffix = "" if number == 1 else f"-{number}"
        candidate = f"{base[:50-len(suffix)]}{suffix}"
        if not db.session.query(model.id).filter(model.code == candidate).first():
            return candidate


def _supplier(name):
    name = _text(name)
    supplier = Supplier.query.filter(db.func.lower(db.func.trim(Supplier.name)) == name.casefold()).first()
    if supplier:
        return supplier, False
    supplier = Supplier(code=_unique_code(Supplier, "IMP-SUP", name), name=name,
                        default_credit_days=30, active=True)
    db.session.add(supplier)
    db.session.flush()
    return supplier, True


def _item(name, gst_percent, user):
    name = _text(name)
    item = Item.query.filter(db.func.lower(db.func.trim(Item.name)) == name.casefold()).first()
    if item:
        return item, False
    item = Item(code=_unique_code(Item, "IMP-ITEM", name), name=name, unit="pcs",
                gst_percent=gst_percent, minimum_stock=Decimal("0.000"),
                active=True, created_by_id=getattr(user, "id", None))
    db.session.add(item)
    db.session.flush()
    return item, True


def _due_date(bill_date, terms):
    match = re.search(r"(\d+)\s*DAY", terms or "", re.IGNORECASE)
    return bill_date + timedelta(days=int(match.group(1))) if match else None


def import_purchase_workbook(preview, company_id, stock_book_id, user):
    company = db.session.get(Company, int(company_id or 0))
    book = db.session.get(StockBook, int(stock_book_id or 0))
    if not company or not company.active:
        raise ValueError("Select an active company.")
    if not book or not book.active or book.company_id != company.id:
        raise ValueError("Select an active stock book belonging to the selected company.")
    if book.book_type != "GST":
        raise ValueError("This purchase register contains GST columns, so select a GST stock book.")
    result = {"created": [], "skipped": [], "created_suppliers": 0, "created_items": 0}
    for data in preview["purchases"]:
        number = data["bill_number"]
        if Purchase.query.filter_by(company_id=company.id, bill_number=number).first():
            result["skipped"].append({"bill_number": number, "reason": "Bill number already exists."})
            continue
        invalid = [line["item"] for line in data["lines"]
                   if line["quantity"] <= 0 or line["rate"] <= 0]
        if invalid:
            result["skipped"].append({"bill_number": number,
                                      "reason": "Missing quantity or rate for: " + ", ".join(invalid)})
            continue
        if not data["date"] or not data["lines"]:
            result["skipped"].append({"bill_number": number,
                                      "reason": "Missing date or product lines."})
            continue
        try:
            with db.session.begin_nested():
                supplier, supplier_created = _supplier(data["supplier"])
                gst_percent = Decimal("18.00") if data["gst_total"] > 0 else Decimal("0.00")
                lines, items_created = [], 0
                for line in data["lines"]:
                    item, created = _item(line["item"], gst_percent, user)
                    items_created += int(created)
                    lines.append({"item_id": item.id, "quantity": line["quantity"],
                                  "rate": line["rate"], "gst_percent": gst_percent})
                remarks = f"Imported from {preview['filename']}"
                if data["payment_terms"]:
                    remarks += f". Payment terms: {data['payment_terms']}"
                if data["round_off"]:
                    remarks += f". Source round-off: {data['round_off']}"
                due = _due_date(data["date"], data["payment_terms"])
                purchase = create_purchase({
                    "company_id": company.id, "stock_book_id": book.id,
                    "supplier_id": supplier.id, "purchase_type": "GST",
                    "bill_number": number, "bill_date": data["date"].isoformat(),
                    "due_date": due.isoformat() if due else "", "remarks": remarks,
                }, lines, user)
            result["created"].append(purchase.bill_number)
            result["created_suppliers"] += int(supplier_created)
            result["created_items"] += items_created
        except Exception as exc:
            result["skipped"].append({"bill_number": number, "reason": str(exc)})
    return result
